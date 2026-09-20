from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(r"D:\CDriveArchive\Desktop-contest\2026年辽宁省大学生物理实验竞赛附件")
OUT = Path("work/extracted_attachment_text.json")


def clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def read_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # pragma: no cover - diagnostic path
            text = f"[page extraction failed: {exc}]"
        pages.append(clean(text))
    return {
        "kind": "pdf",
        "pages": len(reader.pages),
        "text": "\n\n".join(f"--- Page {i} ---\n{page}" for i, page in enumerate(pages, start=1)),
    }


def read_docx(path: Path) -> dict:
    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table_idx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))
        if rows:
            parts.append(f"[Table {table_idx}]\n" + "\n".join(rows))
    return {"kind": "docx", "paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "text": clean("\n".join(parts))}


def read_doc_binary_metadata(path: Path) -> dict:
    # The legacy .doc template is mostly needed as a source/template reference here.
    return {"kind": "doc", "bytes": path.stat().st_size, "text": "[legacy .doc file; not text-extracted]"}


def pptx_slide_text(path: Path) -> dict:
    ns = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    }
    slides: list[tuple[int, str]] = []
    with zipfile.ZipFile(path) as zf:
        names = sorted(
            (name for name in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", name)),
            key=lambda n: int(re.search(r"slide(\d+)\.xml", n).group(1)),
        )
        for idx, name in enumerate(names, start=1):
            root = ET.fromstring(zf.read(name))
            texts = [node.text for node in root.findall(".//a:t", ns) if node.text]
            slides.append((idx, clean("\n".join(texts))))
    return {
        "kind": "pptx",
        "slides": len(slides),
        "text": "\n\n".join(f"--- Slide {idx} ---\n{text}" for idx, text in slides),
    }


def read_xlsx(path: Path) -> dict:
    wb = load_workbook(path, data_only=False, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                rows.append(values)
        preview = rows[:20]
        sheets.append({"name": ws.title, "rows": len(rows), "preview": preview})
    return {"kind": "xlsx", "sheets": sheets, "text": json.dumps(sheets, ensure_ascii=False, indent=2)}


def read_text(path: Path) -> dict:
    return {"kind": path.suffix.lower().lstrip("."), "text": clean(path.read_text(encoding="utf-8", errors="replace"))}


def main() -> int:
    results = {}
    for path in sorted(ROOT.iterdir(), key=lambda p: p.name):
        if path.is_dir():
            continue
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                data = read_pdf(path)
            elif suffix == ".docx":
                data = read_docx(path)
            elif suffix == ".doc":
                data = read_doc_binary_metadata(path)
            elif suffix == ".pptx":
                data = pptx_slide_text(path)
            elif suffix == ".xlsx":
                data = read_xlsx(path)
            elif suffix in {".md", ".txt"}:
                data = read_text(path)
            else:
                data = {"kind": suffix.lstrip(".") or "unknown", "bytes": path.stat().st_size, "text": "[binary or image/video file]"}
        except Exception as exc:
            data = {"kind": suffix.lstrip(".") or "unknown", "error": str(exc), "text": ""}
        data["size"] = path.stat().st_size
        results[path.name] = data

    OUT.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {OUT} with {len(results)} files")
    for name, data in results.items():
        text = data.get("text", "")
        print(f"{name}\t{data.get('kind')}\t{text[:120].replace(chr(10), ' ')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
