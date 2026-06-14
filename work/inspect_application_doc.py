from __future__ import annotations

from pathlib import Path

from docx import Document

DOC = Path(r"D:\CDriveArchive\Desktop-contest\2026年辽宁省大学生物理实验竞赛附件\量子隧穿虚拟仿真实验平台_项目申请书_最终上传版.docx")
doc = Document(str(DOC))

print(f"doc={DOC}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
print("\nPARAGRAPHS")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        print(f"P{i:03d} [{p.style.name}] {text}")

print("\nTABLES")
for ti, table in enumerate(doc.tables):
    print(f"\nTABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
    for ri, row in enumerate(table.rows):
        cells = [" / ".join(x.strip() for x in c.text.splitlines() if x.strip()) for c in row.cells]
        print(f"R{ri:02d}: " + " | ".join(cells))
