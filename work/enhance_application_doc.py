from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


SRC = Path(r"D:\CDriveArchive\Desktop-contest\2026年辽宁省大学生物理实验竞赛附件\量子隧穿虚拟仿真实验平台_项目申请书_最终上传版.docx")
OUT = Path(r"C:\Users\jiang\Documents\Codex\2026-06-14\d-cdrivearchive-desktop-contest-2026\outputs\量子隧穿虚拟仿真实验平台_项目申请书_四项内容强化版.docx")


def set_cjk_font(run, name: str = "宋体", size_pt: float = 10.5, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def format_paragraph(paragraph, first_line: bool = True) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(6)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line:
        fmt.first_line_indent = Pt(21)
    for run in paragraph.runs:
        set_cjk_font(run)


def insert_after(paragraph, text: str, bold_prefix: str | None = None, first_line: bool = True):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    p.style = paragraph.style
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_cjk_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_cjk_font(r2)
    else:
        r = p.add_run(text)
        set_cjk_font(r)
    format_paragraph(p, first_line=first_line)
    return p


def set_heading_like(paragraph) -> None:
    for run in paragraph.runs:
        set_cjk_font(run, name="黑体", size_pt=12, bold=True)
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(8)
    fmt.space_after = Pt(6)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.first_line_indent = None


def find_para(doc: Document, contains: str):
    for p in doc.paragraphs:
        if contains in p.text:
            return p
    raise RuntimeError(f"Paragraph not found: {contains}")


def replace_text_keep_simple(paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    text = paragraph.text.replace(old, new)
    paragraph.clear()
    run = paragraph.add_run(text)
    set_cjk_font(run)
    format_paragraph(paragraph)
    return True


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))

    # Keep the original form, but correct the sampling count so document and code agree.
    for p in doc.paragraphs:
        if "参数扫描在指定范围均匀采样 100 个数据点" in p.text:
            replace_text_keep_simple(
                p,
                "参数扫描在指定范围均匀采样 100 个数据点",
                "参数扫描在指定范围均匀采样 101 个数据点（含区间两端点）",
            )
        if "每类 100 个数据点" in p.text:
            replace_text_keep_simple(
                p,
                "每类 100 个数据点",
                "每类 101 个数据点（含区间两端点）",
            )

    anchor = find_para(doc, "内置 T 范围检查")

    p = insert_after(anchor, "1.5 参数设置范围与教学覆盖面", first_line=False)
    set_heading_like(p)
    texts = [
        "本平台采用自然单位制（hbar = 1），将能量、势垒宽度和质量处理为无量纲可调参数。参数范围既覆盖典型量子隧穿教学区间，又避免极端数值导致图像压缩、双曲函数溢出或学生难以解释的非教学场景。",
        "交互滑块与数值输入的实际范围为：粒子能量 E = 0.05-5.00，步长 0.01；势垒高度 V0 = 0.10-5.00，步长 0.01；势垒宽度 a = 0.10-6.00，步长 0.01；粒子质量 m = 0.20-5.00，步长 0.01。输入框会自动限制在有效范围内，保证所有计算参数为正且处于可解释区间。",
        "参数扫描模块与滑块范围保持一致：T-E 扫描为 E = 0.05-5.00，T-a 扫描为 a = 0.10-6.00，T-V0 扫描为 V0 = 0.10-5.00。每次扫描生成 101 个含端点的均匀采样点，便于观察单一变量对透射概率 T 和反射概率 R 的连续影响。",
        "预设场景覆盖六类教学重点：明显隧穿、弱隧穿、接近势垒顶部、越垒传播、宽势垒抑制和重粒子隧穿。通过这些预设，学生可以从“经典禁止但量子允许”“能量接近势垒时的敏感变化”“能量高于势垒仍存在量子反射”“质量与宽度导致指数抑制”等角度建立完整概念链条。",
    ]
    cur = p
    for t in texts:
        cur = insert_after(cur, t)

    p = insert_after(cur, "1.6 自主构建计算与验证过程", first_line=False)
    set_heading_like(p)
    texts = [
        "计算核心由团队自主实现，集中在源码的 tunneling 计算模块中，而不是调用现成物理仿真黑箱。程序首先根据 E 与 V0 的关系自动分为三种情形：E < V0 的隧穿区、E > V0 的越垒区，以及 E 约等于 V0 的临界区；不同情形分别采用 sinh、sin 和临界极限形式计算透射概率。",
        "在概率计算之外，程序进一步求解三段区域波函数系数。入射区采用 incident + reflected 的形式，势垒区根据物理区间采用指数衰减、振荡或线性极限形式，透射区采用 transmitted 的形式；通过 x = 0 与 x = a 处波函数及一阶导数连续条件，计算反射振幅、透射振幅和势垒内复系数，用于绘制 |psi(x)|^2 概率密度。",
        "验证过程分为四层：第一层为公式比对，透射概率公式与标准量子力学教材中的一维方势垒解析结果一致；第二层为守恒校验，程序实时检查 T、R 是否落在 [0,1]，并验证 |T+R-1| < 10^(-6)；第三层为极限行为校验，包括高能极限 T 接近 1、宽势垒/高势垒极限 T 接近 0、共振条件下 T 接近 1；第四层为趋势校验，E 增大时 T 总体增大，V0、a、m 增大时隧穿区 T 总体降低。",
        "工程验证方面，项目使用 TypeScript 类型约束参数结构和返回结果，构建流程包含类型检查与生产构建；典型场景、参数扫描、导出数据和自动报告功能均围绕同一计算函数生成，减少“图表、表格、报告结果不一致”的风险。",
    ]
    cur = p
    for t in texts:
        cur = insert_after(cur, t)

    anchor = find_para(doc, "运行环境：普通 Windows")
    p = insert_after(anchor, "2.4 每位成员贡献与开发历程", first_line=False)
    set_heading_like(p)
    texts = [
        "团队采用“物理模型-程序实现-教学表达-材料整合”的协作方式推进。按申报书封面成员顺序，第一成员主要负责项目统筹、物理模型推导、核心算法实现、前端架构搭建和最终集成；第二成员主要负责交互界面设计、图表可视化、预设实验场景、参数扫描与数据导出功能；第三成员主要负责文献整理、教学任务设计、模型验证用例、运行说明、展示材料和文档校对。三名成员共同参与功能测试、演示流程排练和校内赛后材料完善。",
        "开发历程可概括为五个阶段。第一阶段完成竞赛方向研读与选题定位，确定面向“教学资源和虚仿”的一维方势垒量子隧穿实验平台；第二阶段完成解析公式整理、边界条件推导和核心计算函数实现；第三阶段完成 React + TypeScript 前端搭建，形成参数控制、波函数可视化、经典-量子对比和实验数据表；第四阶段补充参数扫描、数据导出、自动实验报告、教学说明和典型预设场景；第五阶段围绕校内比赛反馈完善申报书、设计报告、PPT、视频和源码包，进一步检查匿名化、运行说明和材料一致性。",
        "分工并非简单切块，而是围绕同一计算核心迭代协作：算法结果先由成员间交叉核算，再进入可视化页面；页面功能稳定后再提炼为教学文案、视频脚本和报告内容；材料定稿前再次回到程序中核对参数范围、典型数据和导出结果，形成“模型-程序-文档-演示”闭环。",
    ]
    cur = p
    for t in texts:
        cur = insert_after(cur, t)

    anchor = find_para(doc, "数据有效性校验")
    p = insert_after(anchor, "3.4 局限性与改进思路", first_line=False)
    set_heading_like(p)
    texts = [
        "现阶段平台采用一维方势垒、定态散射和自然单位制，突出课堂教学中最核心、最可解释的量子隧穿机制。该处理便于学生理解透射概率、反射概率和势垒参数之间的关系，但尚未覆盖真实器件中可能出现的多维势场、多体相互作用、非弹性散射、退相干和实验噪声等复杂因素。",
        "当前波函数展示侧重空间分布和概率密度，尚未加入含时波包演化。因此学生能够清晰观察定态结果，却还不能直接看到波包从入射、进入势垒到透射/反射的动态过程。后续可采用 Crank-Nicolson 等数值方法求解含时薛定谔方程，增加动画演化、时间步长调节和动态概率流显示。",
        "当前参数使用无量纲自然单位，适合概念教学和规律观察。后续可增加 eV、nm、fs 等真实单位换算，引入电子质量、STM 隧穿距离、半导体势垒等案例，使平台从“概念演示”进一步扩展到“近真实实验参数分析”。",
        "模型扩展方向包括双势垒/共振隧穿二极管、任意形状势垒、WKB 近似与数值传输矩阵方法对比、真实 STM 数据导入与拟合、学生答题反馈和学习记录模块。这样可以在保持现有低成本、易部署优势的基础上，提高平台的研究性、开放性和教学评价能力。",
    ]
    cur = p
    for t in texts:
        cur = insert_after(cur, t)

    # Add a small reinforcing sentence to the conclusion without changing its role.
    conclusion = find_para(doc, "作品解决了量子隧穿概念抽象")
    conclusion.add_run(
        " 本次强化后，申报书进一步明确了团队分工、真实参数范围、计算验证链条和后续迭代路线，使作品的自主开发过程与可持续改进价值更加清晰。"
    )
    for run in conclusion.runs:
        set_cjk_font(run)
    format_paragraph(conclusion)

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
