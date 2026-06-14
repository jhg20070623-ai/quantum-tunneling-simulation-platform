from __future__ import annotations

from pathlib import Path

from docx import Document

DOC = Path(r"C:\Users\jiang\Documents\Codex\2026-06-14\d-cdrivearchive-desktop-contest-2026\outputs\量子隧穿虚拟仿真实验平台_项目申请书_四项内容强化版.docx")
doc = Document(str(DOC))
text = "\n".join(p.text for p in doc.paragraphs)
checks = [
    "每位成员贡献与开发历程",
    "参数设置范围与教学覆盖面",
    "局限性与改进思路",
    "自主构建计算与验证过程",
    "E = 0.05-5.00",
    "V0 = 0.10-5.00",
    "a = 0.10-6.00",
    "m = 0.20-5.00",
    "|T+R-1| < 10^(-6)",
    "Crank-Nicolson",
]
print(f"doc={DOC}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
for item in checks:
    print(f"{item}: {'OK' if item in text else 'MISSING'}")
